import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.ttk import Combobox
from openpyxl import Workbook, load_workbook
import csv
import os
import json
from cryptography.fernet import Fernet
import sys
sys.path.append(r"C:\Users\nicol")  # Add the path to CLIPYTHONPY.py dynamically
from CLIPYTHONPY import generate_schedule, save_schedule, load_schedule, delete_schedule


class StudyPlannerGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Study Planner GUI")
        self.root.geometry("1000x700")
        self.root.configure(bg="#f0f8ff")  # Set background color
        self.root.resizable(True, True)

        self.style = ttk.Style()
        self.style.configure("TButton", font=("Helvetica", 10), padding=5)
        self.style.configure("TLabel", font=("Helvetica", 12), background="#f0f8ff")

        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill=tk.BOTH, expand=True)

        self.create_main_tab()
        self.create_saved_schedules_tab()
        self.create_statistics_tab()

        self.schedule = None
        self.log_file = "statistics_log.txt"

        # Bind the close event to prompt for saving
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def create_main_tab(self):
        main_frame = ttk.Frame(self.notebook)
        self.notebook.add(main_frame, text="Main")

        ttk.Label(main_frame, text="Enter subjects (comma-separated):").pack(pady=5)
        self.subjects_entry = ttk.Entry(main_frame)
        self.subjects_entry.pack(pady=5)

        ttk.Label(main_frame, text="Enter total study hours per day:").pack(pady=5)
        self.hours_entry = ttk.Entry(main_frame)
        self.hours_entry.pack(pady=5)

        ttk.Label(main_frame, text="Enter number of days:").pack(pady=5)
        self.days_entry = ttk.Entry(main_frame)
        self.days_entry.pack(pady=5)

        ttk.Button(main_frame, text="Generate Schedule", command=self.generate_schedule).pack(pady=10)
        ttk.Button(main_frame, text="Save Schedule", command=self.save_schedule).pack(pady=10)
        ttk.Button(main_frame, text="Load Schedule", command=self.load_schedule).pack(pady=10)

        self.text_frame = ttk.Frame(main_frame)
        self.text_frame.pack(fill=tk.BOTH, expand=True)
        self.scrollbar = tk.Scrollbar(self.text_frame)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.schedule_display = tk.Text(self.text_frame, height=20, width=90, yscrollcommand=self.scrollbar.set)
        self.schedule_display.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.config(command=self.schedule_display.yview)

    def create_saved_schedules_tab(self):
        saved_frame = ttk.Frame(self.notebook)
        self.notebook.add(saved_frame, text="Saved Schedules")

        ttk.Label(saved_frame, text="View and manage your saved schedules here.").pack(pady=10)
        self.saved_schedules_list = tk.Listbox(saved_frame, height=15, width=80)
        self.saved_schedules_list.pack(pady=5)

        ttk.Button(saved_frame, text="Refresh List", command=self.list_saved_schedules).pack(pady=5)
        ttk.Button(saved_frame, text="Load Selected", command=self.load_selected_schedule).pack(pady=5)
        ttk.Button(saved_frame, text="Delete Selected", command=self.delete_selected_schedule).pack(pady=5)

    def create_statistics_tab(self):
        stats_frame = ttk.Frame(self.notebook)
        self.notebook.add(stats_frame, text="Statistics")

        ttk.Label(stats_frame, text="Download action log:").pack(pady=10)
        ttk.Button(stats_frame, text="Download Log", command=self.download_log).pack(pady=5)

    def generate_schedule(self):
        print("[DEBUG] Generating schedule")
        subjects = self.subjects_entry.get().split(',')
        subjects = [s.strip() for s in subjects if s.strip()]
        subjects = list(dict.fromkeys(subjects))
        print(f"[DEBUG] Subjects: {subjects}")
        if not subjects:
            messagebox.showerror("Input Error", "Please enter at least one subject.")
            return
        try:
            total_hours = float(self.hours_entry.get())
            num_days = int(self.days_entry.get())
            print(f"[DEBUG] Total hours: {total_hours}, Number of days: {num_days}")
            if total_hours <= 0 or num_days <= 0:
                raise ValueError
            self.schedule = generate_schedule(subjects, total_hours, num_days)
            self.display_schedule()
            self.log_action("Generated schedule")
        except ValueError:
            messagebox.showerror("Input Error", "Please enter valid positive numbers for hours and days.")

    def display_schedule(self):
        print("[DEBUG] Displaying schedule")
        self.schedule_display.delete(1.0, tk.END)
        if self.schedule:
            for day in self.schedule:
                self.schedule_display.insert(tk.END, f"Day {day['Day']}\n")
                for subject in day['Subjects']:
                    self.schedule_display.insert(tk.END, f"  {subject['Subject']}: {subject['Hours']} hours\n")
                self.schedule_display.insert(tk.END, "\n")

    def save_schedule(self):
        print("[DEBUG] Saving schedule")
        if not self.schedule:
            messagebox.showerror("Save Error", "No schedule to save. Please generate one first.")
            return
        file_path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[
            ("JSON files", "*.json"),
            ("Excel files", "*.xlsx"),
            ("CSV files", "*.csv")
        ])
        if not file_path:
            return
        ext = os.path.splitext(file_path)[1].lower()
        print(f"[DEBUG] File path: {file_path}, Extension: {ext}")
        try:
            if ext == ".json":
                save_schedule(self.schedule, file_path, format="json")
            elif ext == ".xlsx":
                save_schedule(self.schedule, file_path, format="xlsx")
            elif ext == ".csv":
                save_schedule(self.schedule, file_path, format="csv")
            else:
                messagebox.showerror("Save Error", "Unsupported file format.")
            self.log_action(f"Saved schedule to {file_path}")
        except (IOError, OSError) as e:
            print(f"[ERROR] Save failed: {e}")
            messagebox.showerror("Save Error", f"An error occurred while saving the file: {e}")

    def load_schedule(self):
        print("[DEBUG] Loading schedule")
        file_path = filedialog.askopenfilename(filetypes=[
            ("All Supported Files", "*.json *.xlsx *.csv"),
            ("JSON files", "*.json"),
            ("Excel files", "*.xlsx"),
            ("CSV files", "*.csv")
        ])
        if not file_path:
            return
        print(f"[DEBUG] File selected for loading: {file_path}")
        try:
            self.schedule = load_schedule(file_path)
            if self.schedule:
                self.display_schedule()
                self.log_action(f"Loaded schedule from {file_path}")
            else:
                messagebox.showerror("Load Error", "Failed to load the schedule.")
        except Exception as e:
            print(f"[ERROR] Load failed: {e}")
            messagebox.showerror("Load Error", f"An error occurred while loading the schedule: {e}")

    def list_saved_schedules(self):
        print("[DEBUG] Listing saved schedules")
        self.saved_schedules_list.delete(0, tk.END)
        files = [f for f in os.listdir('.') if f.endswith(('.json', '.xlsx', '.csv'))]
        print(f"[DEBUG] Found files: {files}")
        for file in files:
            self.saved_schedules_list.insert(tk.END, file)

    def load_selected_schedule(self):
        print("[DEBUG] Loading selected schedule")
        selected = self.saved_schedules_list.curselection()
        if not selected:
            messagebox.showerror("Load Error", "Please select a schedule to load.")
            return
        file_path = self.saved_schedules_list.get(selected[0])
        print(f"[DEBUG] Selected file: {file_path}")
        file_path = os.path.abspath(file_path)
        try:
            self.schedule = load_schedule(file_path)
            if self.schedule:
                self.display_schedule()
                self.log_action(f"Loaded selected schedule from {file_path}")
            else:
                messagebox.showerror("Load Error", "Failed to load the selected schedule.")
        except Exception as e:
            print(f"[ERROR] Load selected failed: {e}")
            messagebox.showerror("Load Error", f"An error occurred while loading the selected schedule: {e}")

    def delete_selected_schedule(self):
        print("[DEBUG] Deleting selected schedule")
        selected = self.saved_schedules_list.curselection()
        if not selected:
            messagebox.showerror("Delete Error", "Please select a schedule to delete.")
            return
        file_path = self.saved_schedules_list.get(selected[0])
        try:
            os.remove(file_path)
            messagebox.showinfo("Delete Success", f"Schedule '{file_path}' deleted successfully.")
            print(f"[DEBUG] Deleted file: {file_path}")
            self.log_action(f"Deleted schedule {file_path}")
            self.list_saved_schedules()
        except Exception as e:
                        print(f"[ERROR] Delete failed: {e}")
                        messagebox.showerror("Delete Error", f"An error occurred while deleting the schedule: {e}")

    def log_action(self, action):
        """ Log user actions to a text file """
        try:
            with open(self.log_file, 'a') as log:
                log.write(f"{action}\n")
            print(f"[DEBUG] Logged action: {action}")
        except Exception as e:
            print(f"[ERROR] Failed to log action: {e}")

    def download_log(self):
        """ Allow user to download the log file """
        try:
            file_path = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Text files", "*.txt")])
            if not file_path:
                return
            with open(self.log_file, 'r') as log:
                data = log.read()
            with open(file_path, 'w') as file:
                file.write(data)
            messagebox.showinfo("Download Log", f"Log file saved to {file_path}")
            print(f"[DEBUG] Log file downloaded to: {file_path}")
        except FileNotFoundError:
            messagebox.showerror("Download Log", "Log file not found.")
        except Exception as e:
            print(f"[ERROR] Failed to download log: {e}")
            messagebox.showerror("Download Log", f"An error occurred while downloading the log: {e}")

    def on_close(self):
        """ Prompt user to save before exiting """
        if messagebox.askyesno("Exit", "Do you want to save the log before exiting?"):
            self.download_log()
        self.root.destroy()


if __name__ == "__main__":
    root = tk.Tk()
    app = StudyPlannerGUI(root)
    root.mainloop()
